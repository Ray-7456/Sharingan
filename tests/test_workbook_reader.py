"""模板解析测试：用合成的模板夹具验证提取与交叉核对。

夹具刻意复刻真实模板的三处特征：

1. T 列阈值"表头 20 / 条件格式 25"两处不一致；
2. K 列最大值统计窗口只到第 599 行，其余列到 1029 行；
3. Y 列公式含硬编码常量 ``-4.1``。

需要 openpyxl（可选依赖）；未安装时整个测试类跳过。
"""

import json
import tempfile
import unittest
from datetime import datetime, timedelta
from pathlib import Path

from sharingan.workbook import OPENPYXL_AVAILABLE, compare, compare_with_rules
from sharingan.workbook.reader import read_template

if OPENPYXL_AVAILABLE:
    import openpyxl
    from openpyxl.formatting.rule import CellIsRule

# (表头, 英文标签, 第 3 行统计窗口或 None)
BASE_COLUMNS = (
    ("时间", "Time", None),
    ("齿轮箱输入轴轴温(85)", "tag_b", (4, 1029)),
    ("齿轮箱输出轴轴温(85)", "tag_c", (4, 1029)),
    ("齿轮箱油槽温度(75)", "tag_d", (4, 1029)),
    ("发电机定子绕组温度U1(135)", "tag_e", (4, 1029)),
    ("发电机定子绕组温度V1(135)", "tag_f", (4, 1029)),
    ("发电机定子绕组温度W1(135)", "tag_g", (4, 1029)),
    ("发电机轴承A温度(80)", "tag_h", (4, 1029)),
    ("发电机轴承B温度(80)", "tag_i", (4, 1029)),
    ("机舱温度(50)", "tag_j", (4, 1029)),
    ("功率", "tag_k", (4, 599)),  # 窗口与其它列不同（疑似笔误）
    ("塔底柜温度(50)", "tag_l", (4, 1029)),
    ("机舱柜温度(50)", "tag_m", (4, 1029)),
    ("主轴叶轮侧温度(55)", "tag_n", (4, 1029)),
    ("主轴齿箱侧温度(60)", "tag_o", (4, 1029)),
)
DERIVED_COLUMNS = (
    ("齿轮箱前轴与油温差值（25）", "=ABS(B4-D4)"),
    ("齿轮箱后轴与油温差值（20）", "=ABS(C4-D4)"),
    ("齿轮箱前后轴差值（20）", "=ABS(B4-C4)"),
    ("主轴前后差值（25）", "=ABS(N4-O4)"),
    ("发电机前后差值（20）", "=ABS(H4-I4)"),
    ("发电机绕组差值1（15）", "=ABS(E4-F4)"),
    ("发电机绕组差值2（15）", "=ABS(E4-G4)"),
    ("发电机绕组差值3（15）", "=ABS(F4-G4)"),
    ("主轴浮动轴承温度与机舱温度差值（46）", "=ABS(N4-J4)"),
    ("主轴止推轴承温度与机舱温度差值（46）", "=(ABS(O4-J4))-4.1000"),
)
# 条件格式：与表头一致，但 T 列刻意写成 25（真实模板就是这样）
CF_RANGES = ("B3", "C3", "D3", "E3:G3", "H3:I3", "J3", "L3:M3", "N3", "O3",
             "P3", "Q3", "R3", "S3:T3", "U3:W3", "X3:Y3")
CF_THRESHOLDS = {
    "B3": 85, "C3": 85, "D3": 75, "E3:G3": 135, "H3:I3": 80, "J3": 50,
    "L3:M3": 50, "N3": 55, "O3": 60, "P3": 25, "Q3": 20, "R3": 20,
    "S3:T3": 25, "U3:W3": 15, "X3:Y3": 46,
}


def build_template(path: Path, *, data_rows: int = 1026) -> None:
    """按真实模板结构生成一份合成模板（不含任何真实数据）。"""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"

    for index, (header, tag, window) in enumerate(BASE_COLUMNS, start=1):
        sheet.cell(1, index, header)
        sheet.cell(2, index, tag)
        if window:
            letter = openpyxl.utils.get_column_letter(index)
            sheet.cell(3, index, f"=MAX({letter}${window[0]}:{letter}${window[1]})")
    sheet.cell(3, 11, f"=MAX(K4:K{BASE_COLUMNS[10][2][1]})")  # K 列：窗口写法与其它列不同
    sheet.cell(3, 1, "最大值")

    for offset, (header, formula) in enumerate(DERIVED_COLUMNS):
        index = len(BASE_COLUMNS) + 1 + offset
        sheet.cell(1, index, header)
        sheet.cell(2, index, header)
        sheet.cell(4, index, formula)

    for row in range(4, 4 + data_rows):
        sheet.cell(row, 1, datetime(2026, 5, 29) + (row - 4) * timedelta(minutes=10))

    for cell_range, threshold in CF_THRESHOLDS.items():
        sheet.conditional_formatting.add(
            cell_range,
            CellIsRule(
                operator="greaterThanOrEqual", formula=[str(threshold)], stopIfTrue=False
            ),
        )
    workbook.save(path)


VALID_RULES = {
    "version": 1,
    "base_columns": [openpyxl.utils.get_column_letter(index) for index in range(1, 16)],
    "columns": {
        "B": {"limit": 85, "source": "测试"},
        "D": {"limit": 75, "source": "测试"},
        "N": {"limit": 55, "source": "测试"},
    },
    "derived": {
        "P": {"formula": "abs(B-D)", "limit": 25, "source": "测试"},
        "T": {"formula": "abs(H-I)", "limit": 20, "source": "测试：以表头为准"},
        "Y": {"formula": "abs(O-J)-4.1", "limit": 46, "source": "测试"},
    },
    "unmatched": {"mode": "halt"},
}


@unittest.skipUnless(OPENPYXL_AVAILABLE, "未安装 openpyxl")
class ReaderTest(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.path = Path(self._tmp.name) / "template.xlsx"
        build_template(self.path)
        self.profile = read_template(self.path)

    def test_structure(self):
        self.assertEqual(self.profile.sheet, "Sheet1")
        self.assertEqual(self.profile.first_data_row, 4)
        self.assertEqual(self.profile.last_data_row, 1029)  # 1026 行数据 → 第 4..1029 行
        self.assertEqual(self.profile.base_letters[:3], ["A", "B", "C"])
        self.assertIn("P", self.profile.derived_letters)
        self.assertIn("T", self.profile.derived_letters)

    def test_header_and_conditional_limits(self):
        column_b = self.profile.column("B")
        self.assertEqual(column_b.header_limit, 85)
        self.assertEqual(column_b.cf_limits, (85,))
        self.assertEqual(column_b.limit, 85)

        column_t = self.profile.column("T")
        self.assertEqual(column_t.header_limit, 20)
        self.assertEqual(column_t.cf_limits, (25,))  # 冲突来源
        self.assertEqual(column_t.limit, 20)  # 有效值优先表头

    def test_derived_expressions_are_normalized(self):
        self.assertEqual(self.profile.column("P").expression, "abs(B-D)")
        self.assertEqual(self.profile.column("Y").expression, "abs(O-J)-4.1")
        self.assertEqual(self.profile.column("Y").constants, (4.1,))
        # 原始列第 4 行是数据而不是公式
        self.assertIsNone(self.profile.column("B").expression)

    def test_windows(self):
        self.assertEqual(self.profile.column("B").window, (4, 1029))
        self.assertEqual(self.profile.column("K").window, (4, 599))


@unittest.skipUnless(OPENPYXL_AVAILABLE, "未安装 openpyxl")
class FindingsTest(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.directory = Path(self._tmp.name)
        self.path = self.directory / "template.xlsx"
        build_template(self.path)
        self.profile = read_template(self.path)

    def _messages(self, findings, level=None):
        return [f"{f.subject}｜{f.message}" for f in findings if level is None or f.level == level]

    def test_template_self_check_finds_three_known_issues(self):
        findings = compare(self.profile)
        errors = self._messages(findings, "error")
        warns = self._messages(findings, "warn")
        self.assertEqual(len(errors), 1, errors)  # 只有 T 列口径冲突
        self.assertIn("T 列阈值口径冲突", errors[0])
        self.assertIn("20", errors[0])
        self.assertIn("25", errors[0])
        self.assertTrue(any("统计窗口不一致" in message and "K 列" in message for message in warns))
        self.assertTrue(any("硬编码常量" in message and "4.1" in message for message in warns))

    def test_window_covering_data_is_flagged(self):
        profile = read_template(self.path)
        self.assertIsNone(
            next(
                (f for f in compare(profile) if "盖不住" in f.subject),
                None,
            ),
            "数据行数等于窗口末端时不应报警",
        )
        longer = self.directory / "longer.xlsx"
        build_template(longer, data_rows=1062)  # 数据到第 1065 行 > 窗口 1029
        findings = compare(read_template(longer))
        self.assertTrue(any("盖不住" in f.subject for f in findings))

    def test_compare_with_rules_reports_only_real_differences(self):
        findings = compare_with_rules(self.profile, VALID_RULES)
        errors = self._messages(findings, "error")
        # 规则与模板（按表头口径）是一致的，除 T 列冲突外不应有其它冲突
        self.assertEqual(errors, [], errors)

    def test_compare_with_rules_flags_limit_mismatch(self):
        rules = json.loads(json.dumps(VALID_RULES))
        rules["columns"]["B"]["limit"] = 86
        findings = compare_with_rules(self.profile, rules)
        errors = self._messages(findings, "error")
        self.assertTrue(any("B 列阈值" in message and "85" in message and "86" in message for message in errors))

    def test_compare_with_rules_flags_formula_mismatch(self):
        rules = json.loads(json.dumps(VALID_RULES))
        rules["derived"]["T"]["formula"] = "abs(H-I)+0"
        findings = compare_with_rules(self.profile, rules)
        errors = self._messages(findings, "error")
        self.assertTrue(any("T 派生列公式" in message for message in errors))

    def test_compare_with_rules_flags_unknown_column(self):
        rules = json.loads(json.dumps(VALID_RULES))
        rules["columns"]["Z"] = {"limit": 10, "source": "测试"}
        findings = compare_with_rules(self.profile, rules)
        self.assertTrue(any("Z" in f.subject for f in findings))

    def test_combined_compare_levels(self):
        findings = compare(self.profile, VALID_RULES)
        levels = {f.level for f in findings}
        self.assertIn("error", levels)  # T 列冲突来自模板自检
        self.assertTrue(all(f.level in {"error", "warn", "info"} for f in findings))


if __name__ == "__main__":
    unittest.main()
